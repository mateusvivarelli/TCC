"""Gera Planilhas/planilha_calculos_concentrados.xlsx: os calculos feitos fora dos rasters, para a bacia toda,
com formulas vivas e formatacao minima (cabecalhos claros, linha "como e calculado", larguras, numeros).
Entradas: Dados Iniciais/Dados_finais_histereses_corrigido.xlsx, comparacao_SY_lambda005_woodward.csv,
eventos_processados_manuais.xlsx (Q_pico) e TCC_Texto/tabelas/tab_eventos.md."""
import re, datetime as dt, pandas as pd, openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as L

N = 17
csv = pd.read_csv("Dados Iniciais/comparacao_SY_lambda005_woodward.csv")
src = openpyxl.load_workbook("Dados Iniciais/Dados_finais_histereses_corrigido.xlsx", data_only=True).active
serie = [r[:7] for r in src.iter_rows(min_row=2, values_only=True) if r[4] is not None]
qp_src = openpyxl.load_workbook("Dados Iniciais/eventos_processados_manuais.xlsx", data_only=True).active
hd = [c.value for c in qp_src[1]]
qpico = [r[hd.index("Q_pico")] for r in qp_src.iter_rows(min_row=2, values_only=True)
         if r[0] not in ("1B", None) and r[hd.index("T_R_m3")]]

def dtp(s):
    m = re.match(r"(\d+)/(\d+)/(\d+) (\d+)h", s)
    return dt.datetime(2000 + int(m[3]), int(m[2]), int(m[1]), int(m[4]))
tab = []
for l in open("TCC_Texto/tabelas/tab_eventos.md", encoding="utf-8").read().splitlines()[2:2 + N]:
    c = [x.strip() for x in l.strip("|").split("|")]
    tab.append((int(c[0]), dtp(c[1]), dtp(c[2]), *[float(x.replace(",", ".")) for x in c[4:8]]))

# ---------------- estilo minimo ----------------
B = Font(bold=True)
TIT = Font(bold=True, size=13)
NOTA = Font(italic=True, color="666666")
CAB_FILL = PatternFill("solid", fgColor="DDEBF7")
RASTER_FILL = PatternFill("solid", fgColor="FFF2CC")
RES_FILL = PatternFill("solid", fgColor="E2EFDA")
WRAP = Alignment(wrap_text=True, vertical="top")
FMT = {"t": "#,##0", "t1": "#,##0.0", "m3": "#,##0", "mm": "0.0", "pct": "0.0", "r": "0.00",
       "ln": "0.000", "c": "0.000000", "q": "0.000"}

def header(ws, row, cols, widths=None):
    """cols: lista de (cabecalho, como e calculado)."""
    for j, (h, como) in enumerate(cols, start=1):
        a = ws.cell(row, j, h); a.font = B; a.fill = CAB_FILL; a.alignment = WRAP
        b = ws.cell(row + 1, j, como); b.font = NOTA; b.alignment = WRAP
        b.data_type = "s"  # texto que comeca com "=" nao e formula
        if widths:
            ws.column_dimensions[L(j)].width = widths[j - 1] if isinstance(widths, list) else widths
    ws.row_dimensions[row].height = 45
    ws.row_dimensions[row + 1].height = 45

def titulo(ws, texto, explica):
    ws["A1"] = texto; ws["A1"].font = TIT
    ws["A2"] = explica; ws["A2"].font = NOTA
    ws["A2"].alignment = Alignment(wrap_text=False)

def fmt_col(ws, col, r0, r1, f):
    for r in range(r0, r1 + 1):
        ws.cell(r, col).number_format = FMT[f]

wb = Workbook()

# ================= Leia-me =================
w0 = wb.active; w0.title = "Leia-me"
w0.column_dimensions["A"].width = 22; w0.column_dimensions["B"].width = 110
linhas = [
    ("Planilha dos cálculos concentrados do TCC (bacia inteira, fora dos rasters)", None),
    ("", None),
    ("O que é", "Tudo o que o modelo calcula depois que os rasters já deram a soma de E_unit de cada evento: a calibração do coeficiente c, a validação cruzada, as métricas e as comparações. Também refaz, a partir da série horária, o volume, a vazão de base e a carga medida de cada evento."),
    ("Como ler", "Cada aba tem: linha 1 = título; linha 2 = o que a aba faz; depois, uma linha de cabeçalhos (em azul) e, logo abaixo, uma linha em cinza dizendo como cada coluna é calculada. Clique numa célula para ver a fórmula."),
    ("Cores", "Azul = cabeçalho. Amarelo = número que vem dos rasters (calculado em Python sobre os mapas; aqui entra como valor). Verde = resultado final da aba."),
    ("", None),
    ("Abas, na ordem", None),
    ("Serie_horaria", "Dados horários da seção de controle (vazão Q e concentração CSS) e a vazão de base: a da planilha original e a recalculada aqui pelo filtro de Eckhardt (alfa = 0,995; BFImax = 0,70)."),
    ("Eventos", "Uma linha por evento: chuva, volume total, vazão de base, escoamento direto medido (V_dir), carga medida (SY_medido), e os números que vêm dos rasters (volume do SCS-CN e soma de E_unit no MC e no MFM)."),
    ("Volume", "Compara o escoamento direto calculado pelo SCS-CN com o medido (viés, erro, NSE do volume)."),
    ("MC", "Modelo Calculado: SY = c * soma(E_unit), com o escoamento direto do SCS-CN (só chuva). Calibra c, faz a validação cruzada e calcula as métricas."),
    ("MFM", "Modelo Forçado ao Medido: igual ao MC, mas com o escoamento direto medido distribuído pelo padrão do SCS-CN."),
    ("Regra_chuva", "Regra simples SY = c * R (R = chuva do evento, mm). Referência para o MC (usa a mesma informação: só chuva)."),
    ("Regra_Vdir", "Regra simples SY = c * V_dir (escoamento direto medido). Referência para o MFM."),
    ("Regra_TR", "Regra simples SY = c * TR (volume total medido, com a vazão de base). Só para comparação."),
    ("IC", "Formulação concentrada da Iniciação Científica com os VALORES ORIGINAIS (alfa 0,277; beta 0,807; sedimento simulado e medido da planilha da IC). Ajuste nos 17 eventos da própria IC, sem validação cruzada; eventos e carga medida não são os mesmos do TCC."),
    ("Resumo", "Tabela final comparando todos os modelos (puxa os resultados das outras abas)."),
    ("", None),
    ("Calibração de c", "c é escolhido para minimizar o erro em escala log: ln c = média de [ln(SY_medido) - ln(x)], em que x é a soma de E_unit (MC, MFM) ou a grandeza da regra simples. Na prática, c é a média geométrica de SY_medido / x."),
    ("Validação cruzada", "Leave-one-out: para prever cada evento, c é recalculado com os outros 16 eventos (sem ele). As métricas principais usam essas 17 previsões."),
    ("Métricas", "Erro mediano (%): mediana de |simulado - medido| / medido. Fator 2: % de eventos com simulado entre metade e o dobro do medido. NSE em log: 1 - soma[(ln sim - ln med)²] / soma[(ln med - média)²]. NSE linear: o mesmo sem log. Spearman: correlação entre as posições (ranking) do simulado e do medido."),
    ("Unidades", "Volumes em m³; chuva em mm; vazão em m³/s; CSS em mg/L; sedimento em t. E_unit é um índice sem unidade física; c converte o índice em t."),
]
for i, (a, b) in enumerate(linhas, start=1):
    w0.cell(i, 1, a).font = TIT if i == 1 else B
    if b:
        c = w0.cell(i, 2, b); c.alignment = WRAP
        w0.row_dimensions[i].height = max(15, 15 * (len(b) // 100 + 1))

# ================= Serie_horaria =================
ws = wb.create_sheet("Serie_horaria")
titulo(ws, "Série horária na seção de controle",
       "Fonte: Dados_finais_histereses_corrigido.xlsx. A vazão de base é recalculada pelo filtro de Eckhardt com os parâmetros das células J4 e J5.")
header(ws, 4, [("Data e hora", "dado"), ("Vazão Q (m³/s)", "dado"), ("Concentração de sedimento CSS (mg/L)", "dado"),
               ("Vazão de base da planilha original (m³/s)", "dado"),
               ("Vazão de base pelo filtro de Eckhardt (m³/s)", "= MIN(Q; ((1-BFImax)*alfa*base anterior + (1-alfa)*BFImax*Q) / (1-alfa*BFImax))"),
               ("Carga de sedimento na hora (t)", "= Q * CSS * 3600 s / 1.000.000 (g -> t)")],
       [18, 12, 16, 18, 22, 18])
ws["I4"], ws["J4"] = "alfa (parâmetro do filtro)", 0.995
ws["I5"], ws["J5"] = "BFImax (parâmetro do filtro)", 0.70
ws["I4"].font = ws["I5"].font = B
ws.column_dimensions["I"].width = 28
R0 = 6
for i, (d, p1, p2, pt, q, css, bt) in enumerate(serie):
    r = R0 + i
    ws.cell(r, 1, d).number_format = "dd/mm/yyyy hh:mm"
    ws.cell(r, 2, q); ws.cell(r, 3, css); ws.cell(r, 4, bt)
    ws.cell(r, 5, f"=D{r}" if i == 0 else
            f"=MIN(B{r},((1-$J$5)*$J$4*E{r-1}+(1-$J$4)*$J$5*B{r})/(1-$J$4*$J$5))")
    ws.cell(r, 6, f"=B{r}*C{r}*3.6/1000")
last = R0 + len(serie) - 1
fmt_col(ws, 5, R0, last, "q"); fmt_col(ws, 6, R0, last, "q")
ws.freeze_panes = "B6"
S = lambda col: f"Serie_horaria!${col}${R0}:${col}${last}"

# ================= Eventos =================
we = wb.create_sheet("Eventos")
titulo(we, "Os 17 eventos: dados medidos e números que vêm dos rasters",
       "Volumes e carga são somados na Série_horaria entre o início e o fim do evento. As colunas em amarelo vêm dos rasters (Python).")
cols = [("Evento", ""), ("Início", "dado"), ("Fim", "dado"),
        ("Chuva no posto DBT5, P1 (mm)", "dado"), ("Chuva no posto Campo Limpo, P2 (mm)", "dado"),
        ("Chuva do evento R (mm)", "= 0,307*P1 + 0,693*P2"),
        ("Intensidade máxima em 60 min RI (mm/h)", "dado"), ("Vazão de pico (m³/s)", "dado"),
        ("Volume total escoado TR (m³)", "= soma de Q * 3600 na janela"),
        ("Volume da vazão de base (m³)", "= soma da base da planilha * 3600 na janela"),
        ("Volume da vazão de base, Eckhardt recalculado (m³)", "= soma da base recalculada * 3600 (só para conferência)"),
        ("Escoamento direto medido V_dir (m³)", "= TR - volume da base"),
        ("Escoamento direto calculado pelo SCS-CN V_SCS (m³)", "RASTERS: soma de Q de todos os pixels"),
        ("Razão SCS-CN / medido", "= V_SCS / V_dir"),
        ("Soma de E_unit no MC (índice)", "RASTERS: soma do índice de erosão, escoamento direto do SCS-CN"),
        ("Soma de E_unit no MFM (índice)", "RASTERS: soma do índice de erosão, escoamento direto medido"),
        ("Sedimento medido SY_medido (t)", "= soma da carga horária na janela")]
header(we, 4, cols, [8, 15, 15, 11, 11, 11, 12, 11, 14, 14, 15, 14, 15, 11, 13, 13, 13])
E0 = 6
for i, (ev, ini, fim, p1, p2, r_, ri) in enumerate(tab):
    r = E0 + i
    we.cell(r, 1, ev)
    we.cell(r, 2, ini).number_format = "dd/mm/yyyy hh:mm"
    we.cell(r, 3, fim).number_format = "dd/mm/yyyy hh:mm"
    we.cell(r, 4, p1); we.cell(r, 5, p2)
    we.cell(r, 6, f"=0.307*D{r}+0.693*E{r}"); we.cell(r, 7, ri); we.cell(r, 8, qpico[i])
    win = f'{S("A")},">="&B{r},{S("A")},"<="&C{r}'
    we.cell(r, 9, f'=SUMIFS({S("B")},{win})*3600')
    we.cell(r, 10, f'=SUMIFS({S("D")},{win})*3600')
    we.cell(r, 11, f'=SUMIFS({S("E")},{win})*3600')
    we.cell(r, 12, f"=I{r}-J{r}")
    we.cell(r, 13, float(csv.V_dir_SCS_m3[i]))
    we.cell(r, 14, f"=M{r}/L{r}")
    we.cell(r, 15, float(csv.soma_E_unit_calc_t[i])); we.cell(r, 16, float(csv.soma_E_unit_forc_t[i]))
    we.cell(r, 17, f'=SUMIFS({S("F")},{win})')
    for col in (13, 15, 16):
        we.cell(r, col).fill = RASTER_FILL
E1 = E0 + N - 1
for col, f in [(6, "mm"), (8, "r"), (9, "m3"), (10, "m3"), (11, "m3"), (12, "m3"), (13, "m3"), (14, "r"),
               (15, "t"), (16, "t"), (17, "t")]:
    fmt_col(we, col, E0, E1, f)
we.cell(E1 + 1, 1, "Total").font = B
for col in (9, 10, 12, 13, 17):
    c = we.cell(E1 + 1, col, f"=SUM({L(col)}{E0}:{L(col)}{E1})"); c.font = B; c.number_format = FMT["m3"]
we.freeze_panes = "B6"
EV = lambda col, r: f"Eventos!{col}{r}"
RNG = lambda col: f"Eventos!${col}${E0}:${col}${E1}"

# ================= Volume =================
wv = wb.create_sheet("Volume")
titulo(wv, "Escoamento direto: SCS-CN contra o medido",
       "Mostra quanto o SCS-CN (CN-I, lambda = 0,05 com S convertido) erra o escoamento direto de cada evento.")
header(wv, 4, [("Evento", ""), ("Escoamento direto medido V_dir (m³)", "da aba Eventos"),
               ("Escoamento direto do SCS-CN V_SCS (m³)", "da aba Eventos (rasters)"),
               ("Razão V_SCS / V_dir", "= V_SCS / V_dir"),
               ("Erro (%)", "= (razão - 1) * 100"), ("Erro absoluto (%)", "= |erro|"),
               ("Dentro de fator 2?", "sim se a razão está entre 0,5 e 2")],
       [8, 16, 16, 12, 10, 10, 11])
for i in range(N):
    r, e = E0 + i, E0 + i
    wv.cell(r, 1, f"={EV('A', e)}"); wv.cell(r, 2, f"={EV('L', e)}"); wv.cell(r, 3, f"={EV('M', e)}")
    wv.cell(r, 4, f"=C{r}/B{r}"); wv.cell(r, 5, f"=(D{r}-1)*100"); wv.cell(r, 6, f"=ABS(E{r})")
    wv.cell(r, 7, f'=IF(AND(D{r}>=0.5,D{r}<=2),"sim","não")')
fmt_col(wv, 2, E0, E1, "m3"); fmt_col(wv, 3, E0, E1, "m3"); fmt_col(wv, 4, E0, E1, "r")
fmt_col(wv, 5, E0, E1, "pct"); fmt_col(wv, 6, E0, E1, "pct")
vres = [("Viés do total (%)", f"=(SUM(C{E0}:C{E1})/SUM(B{E0}:B{E1})-1)*100", "pct"),
        ("Erro médio absoluto por evento (%)", f"=AVERAGE(F{E0}:F{E1})", "pct"),
        ("Erro mediano por evento (%)", f"=MEDIAN(F{E0}:F{E1})", "pct"),
        ("Razão mediana", f"=MEDIAN(D{E0}:D{E1})", "r"),
        ("Eventos dentro de fator 2 (%)", f'=COUNTIF(G{E0}:G{E1},"sim")/{N}*100', "pct"),
        ("NSE do volume", f"=1-SUMXMY2(C{E0}:C{E1},B{E0}:B{E1})/DEVSQ(B{E0}:B{E1})", "r")]
VR = {}
for k, (lab, f, nf) in enumerate(vres):
    r = E1 + 3 + k
    a = wv.cell(r, 1, lab); a.font = B
    wv.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = wv.cell(r, 4, f); c.fill = RES_FILL; c.font = B; c.number_format = FMT[nf]
    VR[lab] = f"Volume!$D${r}"
wv.freeze_panes = "B6"

# ================= abas dos modelos (mesmo desenho) =================
RES = {}  # nome da aba -> dict de celulas de resultado

def aba_modelo(nome, titulo_txt, explica, x_nome, x_col, x_fmt, c_unid):
    w = wb.create_sheet(nome)
    titulo(w, titulo_txt, explica)
    cols = [("Evento", ""),
            ("Sedimento medido (t)", "da aba Eventos"),
            (x_nome, "da aba Eventos"),
            ("ln(medido)", "= LN(medido)"),
            ("ln(x)", "= LN(x)"),
            ("ln(medido) - ln(x)", "= D - E (o c ótimo é a média desta coluna, em exp)"),
            ("Simulado no ajuste (t)", "= c (com os 17 eventos) * x"),
            ("Erro no ajuste (%)", "= (simulado / medido - 1) * 100"),
            ("c sem este evento (validação cruzada)", "= EXP(média de F nos outros 16 eventos)"),
            ("Simulado na validação cruzada (t)", "= c sem este evento * x"),
            ("Erro na validação cruzada (%)", "= (simulado / medido - 1) * 100"),
            ("Erro absoluto (%)", "= |erro na validação cruzada|"),
            ("Razão simulado / medido", "= simulado na validação cruzada / medido"),
            ("Dentro de fator 2?", "sim se a razão está entre 0,5 e 2"),
            ("ln(simulado)", "= LN(simulado na validação cruzada); usado no NSE em log"),
            ("Posição do medido (1 = menor)", "ranking, para o Spearman"),
            ("Posição do simulado (1 = menor)", "ranking, para o Spearman")]
    header(w, 4, cols, [8, 12, 15, 10, 10, 13, 13, 11, 14, 14, 12, 11, 11, 10, 11, 11, 11])
    rc = E1 + 4  # linha do c calibrado (resultado)
    for i in range(N):
        r, e = E0 + i, E0 + i
        w.cell(r, 1, f"={EV('A', e)}"); w.cell(r, 2, f"={EV('Q', e)}"); w.cell(r, 3, f"={EV(x_col, e)}")
        w.cell(r, 4, f"=LN(B{r})"); w.cell(r, 5, f"=LN(C{r})"); w.cell(r, 6, f"=D{r}-E{r}")
        w.cell(r, 7, f"=$D${rc}*C{r}"); w.cell(r, 8, f"=(G{r}/B{r}-1)*100")
        w.cell(r, 9, f"=EXP((SUM($F${E0}:$F${E1})-F{r})/{N-1})")
        w.cell(r, 10, f"=I{r}*C{r}"); w.cell(r, 11, f"=(J{r}/B{r}-1)*100"); w.cell(r, 12, f"=ABS(K{r})")
        w.cell(r, 13, f"=J{r}/B{r}")
        w.cell(r, 14, f'=IF(AND(M{r}>=0.5,M{r}<=2),"sim","não")')
        w.cell(r, 15, f"=LN(J{r})")
        w.cell(r, 16, f"=RANK(B{r},$B${E0}:$B${E1},1)+(COUNTIF($B${E0}:$B${E1},B{r})-1)/2")
        w.cell(r, 17, f"=RANK(J{r},$J${E0}:$J${E1},1)+(COUNTIF($J${E0}:$J${E1},J{r})-1)/2")
    for col, f in [(2, "t"), (3, x_fmt), (4, "ln"), (5, "ln"), (6, "ln"), (7, "t"), (8, "pct"), (9, "c"),
                   (10, "t"), (11, "pct"), (12, "pct"), (13, "r"), (15, "ln")]:
        fmt_col(w, col, E0, E1, f)
    res = [(f"c calibrado com os 17 eventos ({c_unid})", f"=EXP(AVERAGE(F{E0}:F{E1}))", "c"),
           ("Erro mediano na validação cruzada (%)", f"=MEDIAN(L{E0}:L{E1})", "pct"),
           ("Eventos dentro de fator 2 (%)", f'=COUNTIF(N{E0}:N{E1},"sim")/{N}*100', "pct"),
           ("NSE em escala log (validação cruzada)", f"=1-SUMXMY2(O{E0}:O{E1},D{E0}:D{E1})/DEVSQ(D{E0}:D{E1})", "r"),
           ("NSE linear (validação cruzada; secundário)", f"=1-SUMXMY2(J{E0}:J{E1},B{E0}:B{E1})/DEVSQ(B{E0}:B{E1})", "r"),
           ("Spearman (validação cruzada)", f"=CORREL(P{E0}:P{E1},Q{E0}:Q{E1})", "r")]
    cells = {}
    w.cell(rc - 1, 1, "Resultados").font = TIT
    for k, (lab, f, nf) in enumerate(res):
        r = rc + k
        a = w.cell(r, 1, lab); a.font = B
        w.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        c = w.cell(r, 4, f); c.fill = RES_FILL; c.font = B; c.number_format = FMT[nf]
        cells[k] = f"'{nome}'!$D${r}" if " " in nome else f"{nome}!$D${r}"
    w.freeze_panes = "B6"
    RES[nome] = cells

aba_modelo("MC", "Modelo Calculado (MC): SY = c * soma(E_unit), escoamento calculado pelo SCS-CN",
           "x = soma de E_unit sobre os 153.855 pixels, com o escoamento direto do SCS-CN (só chuva). É a previsão.",
           "x = soma de E_unit no MC (índice)", "O", "t", "t por unidade de E_unit")
aba_modelo("MFM", "Modelo Forçado ao Medido (MFM): SY = c * soma(E_unit), escoamento direto medido",
           "x = soma de E_unit com o escoamento direto medido distribuído pelo padrão do SCS-CN. Não é previsão: mostra o desempenho com o escoamento direto medido.",
           "x = soma de E_unit no MFM (índice)", "P", "t", "t por unidade de E_unit")
aba_modelo("Regra_chuva", "Regra simples: SY = c * R (chuva do evento)",
           "Referência para o MC: usa só a chuva, como o MC.",
           "x = chuva do evento R (mm)", "F", "mm", "t por mm")
aba_modelo("Regra_Vdir", "Regra simples: SY = c * V_dir (escoamento direto medido)",
           "Referência para o MFM: usa o mesmo escoamento direto medido, sem a MUSLE.",
           "x = escoamento direto medido V_dir (m³)", "L", "m3", "t por m³")
aba_modelo("Regra_TR", "Regra simples: SY = c * TR (volume total medido)",
           "Só para comparação: usa o volume total, com a vazão de base.",
           "x = volume total TR (m³)", "I", "m3", "t por m³")

# ================= IC (valores originais) =================
# Valores ORIGINAIS da IC (Vivarelli et al., 2025), sem recalibrar: ver Scripts/_ic_original.py.
import sys as _sys
_sys.path.insert(0, "Scripts")
import _ic_original as ICO
wi = wb.create_sheet("IC")
titulo(wi, "IC (valores originais): Y = alfa * (D * qp)^beta * K * LS * C, alfa = 0,277 e beta = 0,807",
       "Valores da planilha da IC (Drive, 'Cópia de Tabela Resumo', abas CLP Eventos e CLP Resultados Calibrados). Ajuste nos 17 eventos da própria IC, sem validação cruzada. A lista de eventos (24/11/2018 no lugar de 17/12/2018) e a carga medida diferem das do TCC: não comparar evento a evento.")
cols = [("Evento da IC", ""), ("Data (início)", "dado da IC"), ("Volume D (m³)", "dado da IC (volume total, com a vazão de base)"),
        ("Vazão de pico qp (m³/s)", "dado da IC"), ("Sedimento medido na IC (t)", "dado da IC"),
        ("Sedimento simulado pela IC (t)", "dado da IC (alfa 0,277; beta 0,807)"),
        ("Erro (%)", "= (simulado / medido - 1) * 100"), ("Erro absoluto (%)", "= |erro|"),
        ("Razão simulado / medido", "= simulado / medido"), ("Dentro de fator 2?", "sim se a razão está entre 0,5 e 2"),
        ("ln(medido)", "= LN(medido)"), ("ln(simulado)", "= LN(simulado)"),
        ("Posição do medido (1 = menor)", "ranking"), ("Posição do simulado (1 = menor)", "ranking")]
header(wi, 4, cols, [8, 12, 14, 12, 13, 13, 10, 10, 11, 10, 10, 10, 11, 11])
for i in range(N):
    r = E0 + i
    wi.cell(r, 1, i + 1); wi.cell(r, 2, ICO.DATA[i]); wi.cell(r, 3, float(ICO.VOL_M3[i])); wi.cell(r, 4, float(ICO.QP_M3S[i]))
    wi.cell(r, 5, float(ICO.Y_OBS[i])); wi.cell(r, 6, float(ICO.Y_SIM_CAL[i]))
    wi.cell(r, 7, f"=(F{r}/E{r}-1)*100"); wi.cell(r, 8, f"=ABS(G{r})"); wi.cell(r, 9, f"=F{r}/E{r}")
    wi.cell(r, 10, f'=IF(AND(I{r}>=0.5,I{r}<=2),"sim","não")')
    wi.cell(r, 11, f"=LN(E{r})"); wi.cell(r, 12, f"=LN(F{r})")
    wi.cell(r, 13, f"=RANK(E{r},$E${E0}:$E${E1},1)+(COUNTIF($E${E0}:$E${E1},E{r})-1)/2")
    wi.cell(r, 14, f"=RANK(F{r},$F${E0}:$F${E1},1)+(COUNTIF($F${E0}:$F${E1},F{r})-1)/2")
for col, f in [(3, "m3"), (4, "t1"), (5, "t1"), (6, "t1"), (7, "pct"), (8, "pct"), (9, "r"), (11, "ln"), (12, "ln")]:
    fmt_col(wi, col, E0, E1, f)
rr = E1 + 3
wi.cell(rr - 1, 1, "Resultados (ajuste nos 17 eventos da IC, sem validação cruzada)").font = TIT
res = [("Erro mediano absoluto (%)", f"=MEDIAN(H{E0}:H{E1})", "pct"),
       ("Eventos dentro de fator 2 (%)", f'=COUNTIF(J{E0}:J{E1},"sim")/{N}*100', "pct"),
       ("NSE em escala log", f"=1-SUMXMY2(L{E0}:L{E1},K{E0}:K{E1})/DEVSQ(K{E0}:K{E1})", "r"),
       ("NSE linear (secundário)", f"=1-SUMXMY2(F{E0}:F{E1},E{E0}:E{E1})/DEVSQ(E{E0}:E{E1})", "r"),
       ("Spearman", f"=CORREL(M{E0}:M{E1},N{E0}:N{E1})", "r"),
       ("Erro médio com sinal (%; positivo = simulado acima)", f"=AVERAGE(G{E0}:G{E1})", "pct"),
       ("Menor erro (%)", f"=MIN(G{E0}:G{E1})", "pct"),
       ("Maior erro (%)", f"=MAX(G{E0}:G{E1})", "pct")]
icc = {}
for k, (lab, f, nf) in enumerate(res):
    r = rr + k
    wi.cell(r, 1, lab).font = B
    wi.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    c = wi.cell(r, 5, f); c.fill = RES_FILL; c.font = B; c.number_format = FMT[nf]
    icc[k] = f"IC!$E${r}"
wi.freeze_panes = "B6"

# ================= Resumo =================
wr = wb.create_sheet("Resumo")
titulo(wr, "Resumo: desempenho na validação cruzada (17 eventos, calibração em log, carga total; IC: ajuste com valores originais)",
       "Todos os números vêm das outras abas. Erro mediano, fator 2, NSE em log e Spearman são as métricas principais; o NSE linear é secundário.")
header(wr, 4, [("Modelo", ""), ("Informação medida usada", ""), ("Coeficiente c", "c com os 17 eventos"),
               ("Erro mediano (%)", "validação cruzada"), ("Dentro de fator 2 (%)", "validação cruzada"),
               ("NSE em log", "validação cruzada"), ("Spearman", "validação cruzada"),
               ("NSE linear (secundário)", "validação cruzada")],
       [34, 34, 14, 12, 12, 11, 11, 12])
mods = [("MC (Modelo Calculado)", "chuva", "MC"), ("Regra SY = c * chuva", "chuva", "Regra_chuva"),
        ("MFM (Modelo Forçado ao Medido)", "chuva e escoamento direto medido", "MFM"),
        ("Regra SY = c * V_dir", "escoamento direto medido", "Regra_Vdir"),
        ("Regra SY = c * TR", "volume total medido", "Regra_TR"),
        ("IC, valores originais (ajuste, sem validação cruzada; eventos da IC)", "volume total e vazão de pico medidos", "IC")]
for i, (nome, info, aba) in enumerate(mods):
    r = 6 + i
    wr.cell(r, 1, nome).font = B; wr.cell(r, 2, info)
    if aba == "IC":
        wr.cell(r, 3, "alfa 0,277; beta 0,807")
        vals = [icc[0], icc[1], icc[2], icc[4], icc[3]]
    else:
        c = RES[aba]
        wr.cell(r, 3, f"={c[0]}").number_format = FMT["c"]
        vals = [c[1], c[2], c[3], c[5], c[4]]
    for j, (v, nf) in enumerate(zip(vals, ["pct", "pct", "r", "r", "r"])):
        cc = wr.cell(r, 4 + j, f"={v}"); cc.number_format = FMT[nf]; cc.fill = RES_FILL
wr.cell(13, 1, "Escoamento direto do SCS-CN contra o medido (aba Volume)").font = TIT
for k, lab in enumerate(["Viés do total (%)", "Erro médio absoluto por evento (%)", "Erro mediano por evento (%)",
                         "Eventos dentro de fator 2 (%)", "NSE do volume"]):
    r = 14 + k
    wr.cell(r, 1, lab).font = B
    c = wr.cell(r, 3, f"={VR[lab]}"); c.fill = RES_FILL
    c.number_format = FMT["r"] if "NSE" in lab else FMT["pct"]
wr.freeze_panes = "A6"

# Resumo como segunda aba, logo depois do Leia-me
wb.move_sheet("Resumo", offset=-(len(wb.sheetnames) - 2))
wb.save("Planilhas/planilha_calculos_concentrados.xlsx")
print("ok", wb.sheetnames)
